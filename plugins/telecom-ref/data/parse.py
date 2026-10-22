"""Parsers for the telecom-ref sources: spec tables, ABNF command formats and
grouped AVP definitions, and RFC text tables. Pure functions of their input.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import docx

# Small parsing helpers


def clean_name(text: str) -> str:
    text = text.replace("\xa0", " ").strip()
    text = re.sub(r"\s*-\s*", "-", text)
    text = re.sub(r"\s+", "-", text) if re.fullmatch(r"[\w\-/ ]+", text) else text
    return re.sub(r"-{2,}", "-", text).strip("-")


def dedupe_cells(cells: list[str]) -> list[str]:
    out: list[str] = []
    for cell in cells:
        if not out or out[-1] != cell:
            out.append(cell)
    return out


def spec_numbers(text: str) -> list[str]:
    return re.findall(r"\b(\d{2}\.\d{3})\b", text.replace("\xa0", " "))


def rfc_numbers(text: str) -> list[int]:
    return [int(n) for n in re.findall(r"RFC\s*(\d{3,4})", text)]


FLAG_LETTERS = re.compile(r"\b([MVP])\b")


def flag_letters(text: str) -> list[str]:
    return sorted(set(FLAG_LETTERS.findall(text.upper())), key="VMP".index)


def norm_type(text: str | None) -> str | None:
    if text is None:
        return None
    text = text.replace("\xa0", " ").strip()
    if not text or text in {"-", "—"}:
        return None
    aliases = {"DiamIdent": "DiameterIdentity", "DiamURI": "DiameterURI", "Unsigned 32": "Unsigned32"}
    text = aliases.get(text, text)
    return re.sub(r"\s+", "", text)


# --------------------------------------------------------------------------
# ABNF (command code formats and grouped AVP definitions)

# Names may hold stray spaces ("[ IDR- Flags ]"); clean_name removes them.
AVP_REF = re.compile(r"^\s*(\d*\s*\*\s*\d*)?\s*([<\[{])\s*([A-Za-z0-9][\w\- ]*?)\s*([>\]}])\s*(;.*)?$")
COMMAND_HEAD = re.compile(
    r"^\s*<?\s*([A-Za-z][\w\- ]*?)\s*>?\s*::=\s*<\s*Diameter[\s\-]*Header\s*:\s*(\d+)\s*((?:,\s*[A-Za-z]+\s*)*)(?:,\s*(\d+)\s*)?>",
    re.I,
)
GROUPED_HEAD = re.compile(r"^\s*<?\s*([A-Za-z0-9][\w\- ]*?)\s*>?\s*::=\s*<\s*AVP[\s\-]*Header\s*:\s*(\d+)\s*(\d+)?\s*>", re.I)


@dataclass
class AvpRef:
    name: str
    kind: str  # fixed | required | optional
    min: int
    max: int | None  # None: unbounded

    def to_json(self, avp_id: str | None) -> dict:
        return {"avp": avp_id, "name": self.name, "kind": self.kind, "min": self.min, "max": self.max}


def parse_avp_ref(line: str) -> AvpRef | None:
    match = AVP_REF.match(line.replace("\xa0", " ").replace("\t", " "))
    if not match:
        return None
    qual, open_, name, close = match.group(1), match.group(2), match.group(3), match.group(4)
    if {"<": ">", "[": "]", "{": "}"}[open_] != close:
        return None
    kind = {"<": "fixed", "{": "required", "[": "optional"}[open_]
    lo, hi = (0 if kind == "optional" else 1), 1
    if qual:
        q = qual.replace(" ", "")
        a, _, b = q.partition("*")
        lo = int(a) if a else (0 if kind == "optional" else 1)
        hi = int(b) if b else None
    return AvpRef(clean_name(name), kind, lo, hi)


@dataclass
class CommandFormat:
    name: str
    code: int
    request: bool
    proxiable: bool
    error: bool
    app: int | None
    avps: list[AvpRef]
    doc: str
    version: str | None
    clause: str | None


@dataclass
class GroupedFormat:
    name: str
    code: int
    vendor: int | None
    avps: list[AvpRef]
    doc: str
    clause: str | None


def parse_abnf(lines: list[tuple[str, str | None]], doc: str, version: str | None, default_app: int | None,
               app_overrides: dict[int, int] | None = None) -> tuple[list[CommandFormat], list[GroupedFormat]]:
    """lines: (text, clause) pairs in document order."""
    commands: list[CommandFormat] = []
    grouped: list[GroupedFormat] = []
    i = 0
    while i < len(lines):
        text, clause = lines[i]
        text = text.replace("\xa0", " ")
        # ABNF heads sometimes share a paragraph with their first AVP lines.
        head_lines = text.split("\n")
        head = head_lines[0]
        cmd = COMMAND_HEAD.match(head)
        grp = None if cmd else GROUPED_HEAD.match(head)
        if not cmd and not grp:
            i += 1
            continue
        body_texts = head_lines[1:]
        rest = head[(cmd or grp).end():].strip()
        if rest:
            body_texts.insert(0, rest)
        j = i + 1
        blanks = 0
        while j < len(lines):
            nxt = lines[j][0].replace("\xa0", " ")
            if not nxt.strip():
                blanks += 1
                if blanks > 1:
                    break
                j += 1
                continue
            parts = nxt.split("\n")
            if all(parse_avp_ref(p) for p in parts if p.strip()):
                body_texts.extend(parts)
                blanks = 0
                j += 1
                continue
            break
        refs = [r for r in (parse_avp_ref(t) for t in body_texts if t.strip()) if r]
        if cmd:
            flags = [f.strip().upper() for f in cmd.group(3).split(",") if f.strip()]
            code = int(cmd.group(2))
            app = int(cmd.group(4)) if cmd.group(4) else (app_overrides or {}).get(code, default_app)
            commands.append(CommandFormat(clean_name(cmd.group(1)), code, "REQ" in flags, "PXY" in flags,
                                          "ERR" in flags, app, refs, doc, version, clause))
        else:
            vendor = int(grp.group(3)) if grp.group(3) else None
            grouped.append(GroupedFormat(clean_name(grp.group(1)), int(grp.group(2)), vendor, refs, doc, clause))
        i = j
    return commands, grouped


# --------------------------------------------------------------------------
# 3GPP .docx


HEADING_NUMBER = re.compile(r"^\s*((?:\d+[a-z]?|[A-Z])(?:\.\d+)*)\s")
TAB_HEADING = re.compile(r"^((?:\d+[a-z]?|[A-Z])(?:\.\d+)+)\t[A-Z0-9]")


def docx_lines(path: Path) -> list[tuple[str, str | None]]:
    """Body paragraphs (including those inside tables) with the clause each sits under."""
    document = docx.Document(str(path))
    out: list[tuple[str, str | None]] = []
    clause: str | None = None
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for element in document.element.body.iter(f"{ns}p"):
        style = element.find(f"{ns}pPr/{ns}pStyle")
        style_id = style.get(f"{ns}val") if style is not None else ""
        text = "".join(t.text or "" for t in element.iter(f"{ns}t"))
        tabs = len(list(element.iter(f"{ns}tab")))
        if tabs and "\t" not in text:
            # Rebuild the text with tabs where Word stored them as elements.
            pieces = []
            for node in element.iter():
                if node.tag == f"{ns}t":
                    pieces.append(node.text or "")
                elif node.tag == f"{ns}tab":
                    pieces.append("\t")
                elif node.tag == f"{ns}br":
                    pieces.append("\n")
            text = "".join(pieces)
        flat = text.replace("\xa0", " ")
        if re.match(r"^(Heading|H)\d", style_id, re.I):
            number = HEADING_NUMBER.match(flat)
            if number:
                clause = number.group(1)
        elif not style_id and (number := TAB_HEADING.match(flat)):
            # Converted .doc files lose heading styles but keep "7.3.2<TAB>Title".
            clause = number.group(1)
        out.append((text, clause))
    return out


def docx_headings(lines: list[tuple[str, str | None]]) -> dict[str, str]:
    """Clause number -> heading title."""
    titles: dict[str, str] = {}
    for text, clause in lines:
        if clause and text.replace("\xa0", " ").strip().startswith(clause):
            titles.setdefault(clause, text.replace("\xa0", " ").strip()[len(clause):].strip())
    return titles


@dataclass
class FlagRow:
    name: str
    code: int | None  # None: an RFC table that lists AVPs by name only
    clause: str | None
    type: str | None
    must: list[str]
    may: list[str]
    should_not: list[str]
    must_not: list[str]
    encr: str | None


def docx_flag_rows(path: Path) -> list[FlagRow]:
    document = docx.Document(str(path))
    rows: list[FlagRow] = []
    for table in document.tables:
        header_rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows[:3]]
        header_text = " ".join(" ".join(r) for r in header_rows)
        if "AVP" not in header_text or "Must" not in header_text and "MUST" not in header_text:
            continue
        # Find the header row that names the columns.
        names_row = None
        for idx, cells in enumerate(header_rows):
            lowered = [c.lower() for c in cells]
            if any("must" == c or c.startswith("must") for c in lowered) and any("code" in c for c in lowered):
                names_row = idx
                break
        if names_row is None:
            continue
        cols = [c.lower() for c in header_rows[names_row]]

        def col(*keys: str, exclude: tuple[str, ...] = ()) -> int | None:
            for k, c in enumerate(cols):
                if any(key in c for key in keys) and not any(x in c for x in exclude):
                    return k
            return None

        i_name = col("attribute", "avp name", "name")
        i_code = col("code")
        i_clause = col("clause", "section", "defined")
        i_type = col("type")
        i_must = next((k for k, c in enumerate(cols) if c == "must"), None)
        i_may = next((k for k, c in enumerate(cols) if c == "may"), None)
        i_should = col("should")
        i_must_not = col("must not", "mustnot")
        i_encr = col("encr")
        if i_name is None or i_code is None or i_must is None:
            continue
        for row in table.rows[names_row + 1:]:
            cells = [c.text.strip() for c in row.cells]

            def cell(k: int | None) -> str:
                return cells[k] if k is not None and k < len(cells) else ""

            code = re.sub(r"\D", "", cell(i_code))
            if not code or not cell(i_name):
                continue
            rows.append(FlagRow(
                name=clean_name(cell(i_name)),
                code=int(code),
                clause=(cell(i_clause).replace("\xa0", " ").strip() or None),
                type=norm_type(cell(i_type)),
                must=flag_letters(cell(i_must)),
                may=flag_letters(cell(i_may)),
                should_not=flag_letters(cell(i_should)),
                must_not=flag_letters(cell(i_must_not)),
                encr=(cell(i_encr) or None),
            ))
    return rows


# --------------------------------------------------------------------------
# RFC text


PAGE_JUNK = re.compile(r"^(RFC \d+ .*\d{4}|.*\[Page \d+\])\s*$")


def rfc_lines(text: str) -> list[tuple[str, str | None]]:
    out: list[tuple[str, str | None]] = []
    clause: str | None = None
    for raw in text.replace("\f", "\n").split("\n"):
        if PAGE_JUNK.match(raw):
            continue
        heading = re.match(r"^(\d+(?:\.\d+)*)\.?\s{1,3}\S", raw)
        if heading:
            clause = heading.group(1)
        out.append((raw, clause))
    # Page breaks leave runs of blank lines inside ABNF blocks; collapse them.
    collapsed: list[tuple[str, str | None]] = []
    for line in out:
        if not line[0].strip() and collapsed and not collapsed[-1][0].strip():
            continue
        collapsed.append(line)
    return collapsed


# Name, optional code, defining section, optional data type, then the flag columns.
RFC_FLAG_ROW = re.compile(
    r"^\s{1,8}\|?([A-Za-z][\w\-]*-?)\s+(?:(\d+)\s+)?(\d+(?:\.\d+)+)\s+(?:([A-Za-z]\w*)\s*)?\|(.*)$"
)


def rfc_flag_rows(lines: list[tuple[str, str | None]]) -> list[FlagRow]:
    rows: list[FlagRow] = []
    header: list[str] = []
    pipe_block: list[list[str]] = []
    for text, _ in lines:
        stripped = text.strip()
        if "-" in stripped and (re.match(r"^-+\|", stripped) or re.fullmatch(r"[-+|\s]+", stripped)):
            # A rule line: the pipe lines above it, read column by column, name the flag columns.
            if pipe_block:
                width = max(len(r) for r in pipe_block)
                header = [" ".join(r[k] for r in pipe_block if k < len(r)).strip().upper() for k in range(width)]
            pipe_block = []
            continue
        if "|" in text and not RFC_FLAG_ROW.match(text):
            cells = [c.strip() for c in text.split("|")[1:]]
            pipe_block.append(cells)
            continue
        pipe_block = [] if "|" not in text else pipe_block
        match = RFC_FLAG_ROW.match(text)
        if not match or not header:
            continue
        cells = [c.strip() for c in match.group(5).split("|")]
        named = dict(zip(header, cells))
        rows.append(FlagRow(
            name=match.group(1),
            code=int(match.group(2)) if match.group(2) else None,
            clause=match.group(3),
            type=norm_type(match.group(4)),
            must=flag_letters(named.get("MUST", "")),
            may=flag_letters(named.get("MAY", "")),
            should_not=flag_letters(named.get("SHLD NOT", "") or named.get("SHOULD NOT", "")),
            must_not=flag_letters(named.get("MUST NOT", "")),
            encr=None,
        ))
    return rows


def rfc_command_abbrevs(lines: list[tuple[str, str | None]]) -> dict[int, list[tuple[str, str, bool]]]:
    """RFC command tables: code -> [(full name, abbrev, is_request)]."""
    out: dict[int, list[tuple[str, str, bool]]] = {}
    pending = ""
    for text, _ in lines:
        match = re.match(r"^\s{2,}([A-Z][\w\-]*-?)\s+([A-Z]{3})\s+(\d+)\s+(\d+(?:\.\d+)*)\s*$", text)
        if match:
            pending = match.group(1)
            name, abbrev, code = match.group(1), match.group(2), int(match.group(3))
            out.setdefault(code, []).append((name, abbrev, name.endswith("Request") or abbrev.endswith("R")))
            continue
        cont = re.match(r"^\s{6,}(Request|Answer)\s*$", text)
        if cont and pending.endswith("-"):
            for code, entries in out.items():
                if entries and entries[-1][0] == pending:
                    full, abbrev, _ = entries[-1]
                    entries[-1] = (full + cont.group(1), abbrev, cont.group(1) == "Request")
            pending = ""
    return out


